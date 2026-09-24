"""Unit and offline integration tests for the lunoviq package (no network)."""
import io
import json

import openpyxl
import pytest

from conftest import ROOT
from lunoviq import http, market_inputs, pipeline
from lunoviq.excel import writer
from lunoviq.providers import rates, yahoo
from lunoviq.providers.sec_edgar import SecEdgarProvider
from lunoviq.schema import AUTO, OVERRIDE, TEMPLATE, DataPoint, PriceHistory

KO = ROOT / "data" / "fixtures" / "ko.json"
TEMPLATE_PATH = ROOT / "Lunoviq_Master_Financial_Model_v2.xlsx"


@pytest.fixture(scope="module")
def ko():
    return SecEdgarProvider(facts_path=KO).financials("KO", 3)


# ---------------------------------------------------------------- schema / SEC
def test_sec_provider_maps_to_schema(ko):
    assert ko.fiscal_years == [2023, 2024, 2025]
    dp = ko.items["revenue"][2025]
    assert dp.value == 47_941_000_000 and dp.source == "SEC EDGAR companyfacts" and dp.method
    assert ko.value("income_tax_total", 2025) == 2_861_000_000
    assert ko.missing() == ["total_liabilities"]
    assert ko.value("operating_income", 2025) is not None
    json.dumps(ko.to_dict())                                  # serialisable for run.json


# ---------------------------------------------------------------- market inputs
def _hist(ticker, rets, start=100.0):
    closes, p = [start], start
    for r in rets:
        p *= 1 + r
        closes.append(p)
    dates = ["%04d-%02d-01" % (2020 + i // 12, i % 12 + 1) for i in range(len(closes))]
    return PriceHistory(ticker, dates, closes, "test")


def test_beta_recovers_known_slope():
    mkt = [0.01 * ((i * 7) % 11 - 5) for i in range(60)]
    stock = [1.5 * r + 0.001 for r in mkt]
    b = market_inputs.beta(_hist("X", stock), _hist("^GSPC", mkt), 60)
    assert b.value == pytest.approx(0.67 * 1.5 + 0.33, abs=1e-4) and b.status == AUTO   # Blume default
    assert "raw 1.5000" in b.method


RATINGS = ([(-100000, 0.199999, "D2/D", 0.19), (3, 4.249999, "A3/A-", 0.0089),
            (4.25, 5.499999, "A2/A", 0.0078), (6.5, 8.499999, "Aa2/AA", 0.0055),
            (8.5, 100000, "Aaa/AAA", 0.004)], "January 2026")


def test_cost_of_debt_synthetic_rating(ko):
    rf = DataPoint(0.05, "ratio", "t", "2026-01-01")
    cov = market_inputs.interest_coverage(ko)
    assert 6.5 < cov < 8.5                                       # KO: EBIT ~13.8bn / interest ~1.65bn
    kd = market_inputs.cost_of_debt(ko, rf, RATINGS)
    assert kd.value == pytest.approx(0.0555) and "Aa2/AA" in kd.method
    assert market_inputs.cost_of_debt(ko, rf, None).value == pytest.approx(0.06)   # flat fallback


def test_effective_tax_rate(ko):
    assert market_inputs.effective_tax_rate(ko).value == pytest.approx(0.1796, abs=1e-4)


def test_ratings_parser(monkeypatch):
    row = "<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td></td><td>x</td></tr>"
    page = "Data used is as of January 2026<table>" + "".join(
        row % (lo, hi, r, sp) for lo, hi, r, sp in
        [(-100000, 0.2, "D2/D", "19.00%"), (0.2, 0.65, "C2/C", "16.00%"), (0.65, 0.8, "Ca2/CC", "12.61%"),
         (0.8, 1.25, "Caa/CCC", "8.85%"), (1.25, 1.5, "B3/B-", "5.09%"), (1.5, 1.75, "B2/B", "3.21%"),
         (1.75, 2, "B1/B+", "2.75%"), (2, 2.25, "Ba2/BB", "1.84%"), (2.25, 2.5, "Ba1/BB+", "1.38%"),
         (2.5, 3, "Baa2/BBB", "1.11%"), (8.5, 100000, "Aaa/AAA", "0.40%")]) + "</table>"
    monkeypatch.setattr(http, "fetch_bytes", lambda *a, **k: (page.encode("latin-1"), 0))
    table, asof = rates.DamodaranRatingsProvider().table()
    assert asof == "January 2026" and table[-1] == (8.5, 100000.0, "Aaa/AAA", 0.004)


def test_capital_weights(ko):
    _, wd, we = market_inputs.capital_weights(ko, DataPoint(88.09, "USD/share", "t", "d"))
    assert wd.value + we.value == pytest.approx(1.0)
    assert wd.value == pytest.approx(45_492e6 / (45_492e6 + 88.09 * 4_313e6), abs=1e-4)


# ---------------------------------------------------------------- provider parsers
def test_treasury_parser(monkeypatch):
    csv_body = b'Date,"1 Mo","10 Yr"\n09/22/2026,4.0,5.05\n09/23/2026,3.99,5.11\n'
    monkeypatch.setattr(http, "fetch_bytes", lambda *a, **k: (csv_body, 0))
    dp = rates.TreasuryProvider().risk_free()
    assert dp.value == pytest.approx(0.0511) and dp.as_of == "2026-09-23"


def test_fred_parser_skips_missing(monkeypatch):
    body = b"observation_date,DGS10\n2026-09-22,5.05\n2026-09-23,.\n"
    monkeypatch.setattr(http, "fetch_bytes", lambda *a, **k: (body, 0))
    assert rates.FredProvider().risk_free().value == pytest.approx(0.0505)


def test_damodaran_parser_tolerates_header_spacing(monkeypatch):
    import datetime as dt
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historical ERP"
    ws.append(["Start of month", "ERP (T12 m with sustainable payout)"])
    ws.append([dt.datetime(2026, 8, 1), 0.0428])
    ws.append([dt.datetime(2026, 9, 1), 0.0414])
    buf = io.BytesIO()
    wb.save(buf)
    monkeypatch.setattr(http, "fetch_bytes", lambda *a, **k: (buf.getvalue(), 0))
    dp = rates.DamodaranProvider().erp()
    assert dp.value == pytest.approx(0.0414) and dp.as_of == "2026-09-01"


def test_yahoo_quote_parser(monkeypatch):
    payload = {"chart": {"result": [{"meta": {"currency": "USD", "regularMarketPrice": 88.09,
                                              "regularMarketTime": 1790193602}}]}}
    monkeypatch.setattr(http, "fetch_json", lambda *a, **k: (payload, 0))
    q = yahoo.YahooProvider().quote("KO")
    assert q.price.value == 88.09 and q.currency == "USD"


# ---------------------------------------------------------------- Excel writer
def test_write_named_fills_whole_range():
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    assert writer.write_named(wb, "drv_TaxRate", 0.2) == [
        "01_Inputs_Historicals!%s52" % c for c in "CDEFG"]
    assert writer.read_named(wb, "drv_TaxRate") == [0.2] * 5


def test_write_named_refuses_formula_cells():
    from openpyxl.workbook.defined_name import DefinedName
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    # 01_Inputs!C98 holds a formula (='00_Dashboard'!$B$12)
    wb.defined_names["tmp_formula"] = DefinedName("tmp_formula", attr_text="'01_Inputs_Historicals'!$C$98")
    with pytest.raises(writer.FormulaProtectedError):
        writer.write_named(wb, "tmp_formula", 1)
    assert wb["01_Inputs_Historicals"]["C98"].value.startswith("=")


# ---------------------------------------------------------------- pipeline (offline)
def _market():
    d = lambda v, u="ratio": DataPoint(v, u, "test", "2026-09-24", "fixture")
    return {"ctl_SharePrice": d(88.09, "USD/share"), "val_RiskFree": d(0.0511), "val_ERP": d(0.0414),
            "val_Beta": d(0.55, "x"), "val_CostOfDebt": d(0.0611), "val_DebtWeight": d(0.1069),
            "val_EquityWeight": d(0.8931), "drv_TaxRate": d(0.1796)}


def test_pipeline_offline_populates_copy(tmp_path):
    before = TEMPLATE_PATH.read_bytes()
    res = pipeline.run("KO", facts_path=KO, recalc=False, out_root=tmp_path, market=_market(),
                       sets=["val_TerminalGrowth=0.02"])
    assert TEMPLATE_PATH.read_bytes() == before
    wb = openpyxl.load_workbook(res["model"])
    assert writer.read_named(wb, "val_Beta") == [0.55]
    assert writer.read_named(wb, "ctl_SharePrice") == [88.09]
    assert writer.read_named(wb, "drv_TaxRate") == [0.1796] * 5
    assert writer.read_named(wb, "val_TerminalGrowth") == [0.02]
    log = {r[0].value: r[4].value for r in wb["09_Sources"].iter_rows(min_row=5)}
    assert log["val_TerminalGrowth"] == OVERRIDE
    assert log["val_ExitMultiple"] == TEMPLATE
    assert log["val_RiskFree"] == AUTO
    rec = json.loads(open(res["run_json"]).read())
    assert rec["inputs"]["val_ERP"]["value"] == 0.0414


@pytest.mark.excel
def test_pipeline_recalc_outputs(tmp_path):
    res = pipeline.run("KO", facts_path=KO, recalc=True, out_root=tmp_path, market=_market())
    out = res["outputs"]
    assert out["health"]["3-Statement"] == "OK"
    assert 0.03 < out["wacc"] < 0.10
    assert out["valuation"]["Current Price"] == pytest.approx(88.09)
