"""Regression tests for templates v5 (net income reconciliation) and v6 (comparables fixes)."""
import openpyxl
import pytest

from conftest import ROOT
from xlsx_inspect import cells, defined_names, diff
import build_master_v5
import build_master_v6

V4, V5, V6 = (ROOT / "templates" / ("Lunoviq_Master_Financial_Model_v%d.xlsx" % n) for n in (4, 5, 6))


@pytest.mark.parametrize("mod,out_name,ref", [(build_master_v5, "v5.xlsx", V5), (build_master_v6, "v6.xlsx", V6)])
def test_builds_reproducible(tmp_path, mod, out_name, ref):
    out = tmp_path / out_name
    mod.build(str(out))
    assert cells(out) == cells(ref)


def test_v5_changes_exactly_the_specified_cells():
    a, b = cells(V4), cells(V5)
    changed = {(s, r) for s in a for r in diff(a[s], b[s])}
    assert changed == set(build_master_v5.cells())
    assert set(defined_names(V5)) - set(defined_names(V4)) == set(build_master_v5.NEW_NAMES)


def test_v6_changes_exactly_the_specified_cells():
    a, b = cells(V5), cells(V6)
    changed = {(s, r) for s in a for r in diff(a[s], b[s])}
    assert changed == set(build_master_v6.cells(a))
    assert b["06_Comparable_Valuation"]["H7"][1].endswith('K7/E7)')          # P/E = market cap / NI


@pytest.mark.excel
def test_v6_without_deals_or_fourth_peer_has_no_errors(tmp_path):
    from recalc import recalc
    from lunoviq.excel.reader import _errors
    wb = openpyxl.load_workbook(V6)
    ws = wb["06_Comparable_Valuation"]
    for r in range(41, 46):
        for col in "ABCDEFGHIJ":
            ws["%s%d" % (col, r)] = None
    for col in "ABCDEIJK":
        ws["%s10" % col] = None
    f = tmp_path / "v6_sparse.xlsx"
    wb.save(f)
    out = openpyxl.load_workbook(recalc(f), data_only=True)
    assert _errors(out) == []
    assert out["00_Dashboard"]["B40"].value == "n/a" and out["00_Dashboard"]["I57"].value == "OK"
    assert isinstance(out["06_Comparable_Valuation"]["H27"].value, float)
