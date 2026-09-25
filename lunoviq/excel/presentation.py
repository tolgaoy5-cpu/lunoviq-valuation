"""
Company-specific content for sheets that are not driven by named ranges:
the trading-comparables peer table, the precedent-transaction block, and
hiding the unit-economics schedules that the Growth-Margin engine ignores.
Every write targets a constant cell; formula cells are refused.
"""
from .. import units
from .writer import FormulaProtectedError

COMP = "06_Comparable_Valuation"
PEER_ROWS = range(7, 11)
DEAL_ROWS = range(41, 46)

# Rows of the unit-economics (volume x price, plant capacity) engine. They hold
# the demo manufacturer's data and do not feed the Growth-Margin model.
UNIT_ECONOMICS_ROWS = {
    "01_Inputs_Historicals": [*range(6, 15), 39, 40, 41, 42, 73, *range(91, 96)],
    "02_Operating_Model": [*range(8, 30), 32, 33, *range(37, 78), 79, *range(231, 241), 245],
}


def _put(ws, ref, value):
    cell = ws[ref]
    if isinstance(cell.value, str) and cell.value.startswith("="):
        raise FormulaProtectedError("%s!%s holds a formula" % (ws.title, ref))
    cell.value = value


def write_peers(wb, peers):
    """Peer table in the model unit: name, EV, revenue, EBITDA, net income, market cap."""
    ws = wb[COMP]
    k = units.scale
    for i, r in enumerate(PEER_ROWS):
        p = peers[i] if i < len(peers) else None
        vals = {"A": None, "B": None, "C": None, "D": None, "E": None, "I": None, "J": None, "K": None}
        if p:
            vals.update(A="%s (%s)" % (p["name"][:26], p["ticker"]), B=k(p["ev"]), C=k(p["revenue"]),
                        D=k(p["ebitda"]), E=k(p["net_income"]), I="US", J=p.get("rationale", ""),
                        K=k(p["market_cap"]))
        for col, v in vals.items():
            _put(ws, "%s%d" % (col, r), v)
    return len(peers)


def tidy(wb):
    """Readability fixes on the generated copy (never on the template)."""
    from copy import copy
    comp = wb[COMP]
    comp.column_dimensions["J"].width = 60           # peer rationale
    comp.column_dimensions["K"].width = 16           # market cap
    inp = wb["01_Inputs_Historicals"]
    english = {
        ("01_Inputs_Historicals", "C102"): ("hacim", "Unit Economics = volume x price (needs internal data). "
                                            "Growth-Margin = revenue growth and cost margins from filings; "
                                            "used for listed companies."),
        ("01_Inputs_Historicals", "C95"): ("Kullanim", "If utilisation exceeds the trigger, a capacity block is "
                                           "added the next year; expansion capex flows into cash flow."),
        ("08_Data_Feed", "A36"): ("denetim", "Reported Totals (audit trail)"),
    }
    for (sheet, ref), (marker, text) in english.items():
        cell = wb[sheet][ref]
        if isinstance(cell.value, str) and marker in cell.value:
            cell.value = text
    # whole-number display: amounts are in USD thousands/millions and periods are counts, so the
    # template's one-decimal formats carry no information; share prices, %, x keep decimals
    whole = {"#,##0.0;[Red](#,##0.0);-": "#,##0;[Red](#,##0);-",
             "#,##0.0;[Red]\\(#,##0.0\\);\\-": "#,##0;[Red]\\(#,##0\\);\\-", "0.0": "0"}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.number_format in whole:
                    c.number_format = whole[c.number_format]
    # DCF: year headers were left-aligned over right-aligned numbers, so each figure
    # appeared to sit under the next year; widen the bridge label column
    from openpyxl.styles import Alignment
    dcf = wb["04_DCF_Valuation"]
    for ref in ("B6", "C6", "D6", "E6", "F6", "G6", "B23", "C23", "D23", "E23", "F23"):
        c = dcf[ref]
        if isinstance(c.value, str):
            al = copy(c.alignment)
            al.horizontal = "right"
            c.alignment = al
    dcf.column_dimensions["D"].width = max(dcf.column_dimensions["D"].width or 10, 24)
    # unit note under each sheet title that lacks one
    from openpyxl.styles import Font
    unit = units.label()
    notes = {"04_DCF_Valuation": "All figures in %s unless stated; per-share values in USD" % unit,
             "05_Sensitivity": "Per-share values in USD; enterprise values in %s" % unit,
             "06_Comparable_Valuation": "All figures in %s unless stated; per-share values in USD; multiples in x" % unit}
    for sheet, text in notes.items():
        ws = wb[sheet]
        if ws["A2"].value is None:
            ws["A2"].value = text
            ws["A2"].font = Font(name="Arial", size=9, italic=True, color="5E6B78")


def clear_precedents(wb):
    """No free source for M&A transaction data: leave the block empty (template
    v6 then shows the method as n/a) rather than showing the demo deals."""
    ws = wb[COMP]
    for r in DEAL_ROWS:
        for col in "ABCDEFGHIJ":
            _put(ws, "%s%d" % (col, r), None)


def hide_unit_economics(wb):
    for sheet, rows in UNIT_ECONOMICS_ROWS.items():
        for r in rows:
            wb[sheet].row_dimensions[r].hidden = True


def print_setup(wb, helper_col=24):
    """Landscape, one page wide, print area = used cells left of the chart-helper
    columns (X onward), so the workbook prints / exports to PDF as readable pages
    instead of hundreds of strips."""
    from openpyxl.utils import get_column_letter
    for ws in wb.worksheets:
        last_col, last_row = 1, 1
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and c.column < helper_col and not ws.row_dimensions[c.row].hidden:
                    last_col, last_row = max(last_col, c.column), max(last_row, c.row)
        ws.print_area = "A1:%s%d" % (get_column_letter(last_col), last_row)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.oddFooter.center.text = "&A  |  page &P of &N"


AN = "07_Analysis_Scenarios"
# 07 benchmark rows -> ratio key (same definitions as the sheet's own rows)
BENCH_ROWS = {25: "ebitda_margin", 26: "debt_ebitda", 27: "roa", 28: "roe",
              112: "gross_margin", 113: "ebitda_margin", 114: "net_margin", 115: "roa", 116: "roe",
              117: "asset_turnover", 118: "ccc", 119: "debt_equity", 120: "debt_ebitda",
              121: "current_ratio", 122: "quick_ratio", 123: "revenue_growth", 124: "ebitda_growth"}


def write_benchmarks(wb, ticker, peers):
    """Peer Low / Median / High from the real peer set (latest fiscal year);
    the template's demo ranges and 'Industry' column are removed."""
    import statistics
    ws = wb[AN]
    for ref in ("B24", "B111"):
        _put(ws, ref, "%s 2030E" % ticker)
    _put(ws, "F111", None)
    filled = 0
    for r, key in BENCH_ROWS.items():
        vals = sorted(p["ratios"][key] for p in peers if p.get("ratios", {}).get(key) is not None)
        lo, med, hi = (vals[0], statistics.median(vals), vals[-1]) if vals else (None, None, None)
        for col, val in zip("CDE", (lo, med, hi)):
            _put(ws, "%s%d" % (col, r), val)
        if r >= 112:
            _put(ws, "F%d" % r, None)
        filled += bool(vals)
    return filled


def relabel_units(wb):
    """Template labels say USD thousands / $000 / 000s; switch them when the model runs in millions."""
    if units.DIV == 1e3:
        return 0
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                new = v.replace("USD thousands", units.label())
                if new.strip() == "$000":
                    new = units.short()
                elif new.strip() == "000s":
                    new = units.shares_label()
                if new != v:
                    c.value = new
                    n += 1
    return n


def compact_style(wb):
    """Denser, analyst-style look: body text in Arial Narrow 9 (as in CFI / Macabacus
    models), headings one step larger, tighter default row height."""
    from copy import copy
    body = {("Arial", 10.0), ("Carlito", 11.0), ("Arial", 9.0)}
    heads = {("Arial", 12.0), ("Carlito", 12.0)}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                key = (c.font.name, c.font.sz)
                if key in body or key in heads:
                    f = copy(c.font)
                    f.name = "Arial Narrow"
                    f.sz = 9.0 if key in body else 11.0
                    c.font = f
        ws.sheet_format.defaultRowHeight = 12.75
        ws.sheet_format.customHeight = True
