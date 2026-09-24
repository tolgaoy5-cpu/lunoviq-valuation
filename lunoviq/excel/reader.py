"""Read calculated outputs back from a recalculated workbook (cached values)."""
import openpyxl

from ..valuation_checks import dcf_checks

HEALTH_ROWS = range(51, 58)          # 00_Dashboard H51:I57 (area, status); I57 = overall
VALUATION_ROWS = range(37, 43)       # 00_Dashboard A37:B42


def outputs(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    d = wb["00_Dashboard"]
    if d["I57"].value is None:
        raise RuntimeError("%s has no calculated values - run the recalculation step" % path)
    return {
        "health": {d["H%d" % r].value: d["I%d" % r].value for r in HEALTH_ROWS},
        "valuation": {d["A%d" % r].value: d["B%d" % r].value for r in VALUATION_ROWS},
        "wacc": wb["04_DCF_Valuation"]["B18"].value,
        "terminal_growth": wb["04_DCF_Valuation"]["B19"].value,
        "exit_multiple": wb["04_DCF_Valuation"]["C49"].value,
        "checks_flagged": _flagged(wb),
        "dcf_checks": dcf_checks(wb),
    }


def _flagged(wb):
    """Every formula cell across the model currently showing CHECK."""
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value == "CHECK":
                    out.append("%s!%s (%s)" % (ws.title, c.coordinate, _label(ws, c)))
    return out


def _label(ws, cell):
    """Nearest text to the left on the same row (dashboards hold several tables per row)."""
    for col in range(cell.column - 1, 0, -1):
        v = ws.cell(row=cell.row, column=col).value
        if isinstance(v, str) and v not in ("OK", "CHECK"):
            return v
    return ""
