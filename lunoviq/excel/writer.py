"""
Excel population through the template's named ranges.

Rules that protect the model:
  * values are written only to cells that currently hold constants; a named
    range that points at a formula raises instead of silently overwriting it
  * the master template is never opened for writing (callers pass a copy)
  * every written value is listed on a '09_Sources' sheet with its source,
    date, method and status, so the workbook carries its own audit trail
"""
from openpyxl.styles import Alignment, Font, PatternFill

from ..schema import MARKET_INPUTS, OVERRIDE, TEMPLATE

LOG_SHEET = "09_Sources"


class FormulaProtectedError(RuntimeError):
    pass


def named_cells(wb, name):
    if name not in wb.defined_names:
        raise KeyError("named range %s not in workbook" % name)
    cells = []
    for sheet, coord in wb.defined_names[name].destinations:
        rng = wb[sheet][coord.replace("$", "")]
        if isinstance(rng, tuple):
            for row in (rng if isinstance(rng[0], tuple) else (rng,)):
                cells.extend(row)
        else:
            cells.append(rng)
    return cells


def read_named(wb, name):
    return [c.value for c in named_cells(wb, name)]


def write_named(wb, name, value):
    """Write one value to every cell of a named range (e.g. a 5-year driver row)."""
    cells = named_cells(wb, name)
    for c in cells:
        if isinstance(c.value, str) and c.value.startswith("="):
            raise FormulaProtectedError("%s -> %s!%s holds a formula" % (name, c.parent.title, c.coordinate))
    for c in cells:
        c.value = value
    return ["%s!%s" % (c.parent.title, c.coordinate) for c in cells]


def apply_inputs(wb, inputs):
    """Write {named_range: DataPoint}; returns log rows."""
    rows = []
    for name, dp in inputs.items():
        if name.startswith("_") or dp.value is None:
            rows.append((name, dp, "(not written)"))
            continue
        where = write_named(wb, name, dp.value)
        rows.append((name, dp, where[0] if len(where) == 1 else "%s:%s" % (where[0], where[-1].split("!")[1])))
    return rows


def template_defaults(wb, names, source="template"):
    """DataPoints for judgement inputs deliberately left at the template value."""
    from ..schema import DataPoint
    out = {}
    for n in names:
        v = read_named(wb, n)[0]
        out[n] = DataPoint(v, "ratio" if isinstance(v, float) and v < 1 else "x", source, "",
                           "left at template value - review", TEMPLATE)
    return out


def write_log(wb, rows, header_note):
    if LOG_SHEET in wb.sheetnames:
        del wb[LOG_SHEET]
    ws = wb.create_sheet(LOG_SHEET)
    ws["A1"] = "Input Sources & Assumptions Log"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = header_note
    heads = ["Named range", "Description", "Value", "Unit", "Status", "Source", "As of", "Method / note", "Cell(s)"]
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    fills = {TEMPLATE: "FFF2CC", OVERRIDE: "DDEBF7"}
    for r, (name, dp, where) in enumerate(rows, 5):
        vals = [name, MARKET_INPUTS.get(name, ""), dp.value, dp.unit, dp.status,
                dp.source, dp.as_of, dp.method, where]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            if dp.status in fills:
                c.fill = PatternFill("solid", fgColor=fills[dp.status])
        ws.cell(row=r, column=8).alignment = Alignment(wrap_text=False)
    for col, w in zip("ABCDEFGHI", (20, 28, 16, 10, 16, 30, 22, 70, 30)):
        ws.column_dimensions[col].width = w
    return ws
